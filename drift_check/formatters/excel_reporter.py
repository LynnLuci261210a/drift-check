"""Excel (XLSX) reporter for drift-check results."""
from __future__ import annotations

import io
from typing import Sequence

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
except ImportError as exc:  # pragma: no cover
    raise ImportError(
        "openpyxl is required for Excel output: pip install openpyxl"
    ) from exc

from drift_check.drift_detector import DriftItem, DriftKind

# Colour palette
_HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
_CHANGED_FILL = PatternFill("solid", fgColor="FFF2CC")
_MISSING_FILL = PatternFill("solid", fgColor="FCE4D6")
_EXTRA_FILL = PatternFill("solid", fgColor="E2EFDA")
_HEADER_FONT = Font(bold=True, color="FFFFFF")

_COLUMNS = ["Resource ID", "Kind", "Attribute", "Expected", "Actual"]


def _kind_label(kind: DriftKind) -> str:
    return {
        DriftKind.CHANGED: "changed",
        DriftKind.MISSING: "missing",
        DriftKind.EXTRA: "extra",
    }[kind]


def _row_fill(kind: DriftKind) -> PatternFill:
    return {
        DriftKind.CHANGED: _CHANGED_FILL,
        DriftKind.MISSING: _MISSING_FILL,
        DriftKind.EXTRA: _EXTRA_FILL,
    }[kind]


def render_excel(items: Sequence[DriftItem]) -> bytes:
    """Return an XLSX workbook as *bytes* representing *items*."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Drift Report"

    # Header row
    for col_idx, header in enumerate(_COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(horizontal="center")

    ws.column_dimensions["A"].width = 36
    ws.column_dimensions["B"].width = 10
    ws.column_dimensions["C"].width = 24
    ws.column_dimensions["D"].width = 30
    ws.column_dimensions["E"].width = 30

    if not items:
        ws.cell(row=2, column=1, value="No drift detected.").font = Font(italic=True)
        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()

    for row_idx, item in enumerate(items, start=2):
        fill = _row_fill(item.kind)
        diffs = item.attributes or {}
        if not diffs:
            diffs = {"": ("", "")}
        for attr, (expected, actual) in diffs.items():
            values = [
                item.resource_id,
                _kind_label(item.kind),
                attr,
                str(expected) if expected is not None else "",
                str(actual) if actual is not None else "",
            ]
            for col_idx, val in enumerate(values, start=1):
                cell = ws.cell(row=row_idx, column=col_idx, value=val)
                cell.fill = fill
            row_idx += 1

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
